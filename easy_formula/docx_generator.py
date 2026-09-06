from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

from .models import FormulaCandidate, RecognitionRecord, VerificationResult

CHINESE_FONT = "Noto Sans CJK SC"

CONF_ZH = {"high": "高", "medium": "中", "low": "低"}
KIND_ZH = {"display": "独立公式", "inline": "行内公式", "visual": "页面视觉识别"}


def _set_run_font(run, western: str = CHINESE_FONT, east_asia: str = CHINESE_FONT, size: float | None = None, bold: bool | None = None):
    run.font.name = western
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = CHINESE_FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = CHINESE_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)

    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)


def _add_label_value(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(label)
    _set_run_font(r1, bold=True)
    r2 = p.add_run(value)
    _set_run_font(r2)


def _add_code_block(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.right_indent = Cm(0.45)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    r = p.add_run(code)
    _set_run_font(r, western="Consolas", east_asia=CHINESE_FONT, size=9.5)


def _add_original_formula(doc: Document, candidate: FormulaCandidate) -> None:
    crop_path = candidate.crop_path
    if not crop_path or not Path(crop_path).exists():
        p = doc.add_paragraph("原始公式图片不可用。")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        with Image.open(crop_path) as img:
            px_width = img.width
            dpi = img.info.get("dpi", (240, 240))[0] or 240
        native_cm = px_width / float(dpi) * 2.54
        if candidate.kind == "inline":
            target_cm = min(7.0, max(2.5, native_cm * 2.0))
        else:
            target_cm = min(14.5, max(5.0, native_cm * 1.15))
        p.add_run().add_picture(crop_path, width=Cm(target_cm))
    except Exception:
        p = doc.add_paragraph("原始公式图片插入失败，请核对工作目录中的公式截图。")


def _ready_latex(candidate: FormulaCandidate, latex: str) -> str:
    if candidate.kind == "inline":
        return f"${latex}$"
    return "\\[\n" + latex + "\n\\]"


def generate_docx(
    source_pdf: str | Path,
    output_docx: str | Path,
    candidates: list[FormulaCandidate],
    records: list[RecognitionRecord],
    verification: dict[str, VerificationResult],
) -> Path:
    source_pdf = Path(source_pdf)
    output_docx = Path(output_docx)
    record_map = {r.candidate_id: r for r in records}

    accepted: list[tuple[FormulaCandidate, RecognitionRecord, VerificationResult]] = []
    for candidate in candidates:
        record = record_map.get(candidate.candidate_id)
        if not record or record.is_formula is not True:
            continue
        vr = verification[candidate.candidate_id]
        accepted.append((candidate, record, vr))

    review = [item for item in accepted if item[2].final_confidence == "low" or item[2].issues]
    display_count = sum(c.kind == "display" for c, _, _ in accepted)
    inline_count = sum(c.kind == "inline" for c, _, _ in accepted)
    visual_count = sum(c.kind == "visual" for c, _, _ in accepted)

    doc = Document()
    _style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("公式提取结果")
    _set_run_font(r, size=20, bold=True)

    _add_label_value(doc, "源文件：", source_pdf.name)
    _add_label_value(doc, "识别公式总数：", str(len(accepted)))
    _add_label_value(doc, "独立公式：", str(display_count))
    _add_label_value(doc, "行内公式：", str(inline_count))
    if visual_count:
        _add_label_value(doc, "页面视觉识别公式：", str(visual_count))
    _add_label_value(doc, "需要人工核对：", str(len(review)))

    doc.add_paragraph()

    for idx, (candidate, record, vr) in enumerate(accepted, start=1):
        heading = doc.add_paragraph(style="Heading 1")
        rr = heading.add_run(f"公式 {idx}")
        _set_run_font(rr, size=15, bold=True)

        _add_label_value(doc, "页码：", str(candidate.page_number))
        _add_label_value(doc, "公式类型：", KIND_ZH.get(candidate.kind, "公式"))
        if candidate.equation_number:
            _add_label_value(doc, "原文公式编号：", candidate.equation_number)

        p = doc.add_paragraph()
        rr = p.add_run("原始公式")
        _set_run_font(rr, bold=True)
        _add_original_formula(doc, candidate)

        p = doc.add_paragraph()
        rr = p.add_run("LaTeX 代码")
        _set_run_font(rr, bold=True)
        _add_code_block(doc, record.latex.strip())

        p = doc.add_paragraph()
        rr = p.add_run("可直接使用的 LaTeX")
        _set_run_font(rr, bold=True)
        _add_code_block(doc, _ready_latex(candidate, record.latex.strip()))

        _add_label_value(doc, "识别置信度：", CONF_ZH.get(vr.final_confidence, "低"))
        if vr.final_confidence == "low":
            p = doc.add_paragraph("此公式可能存在识别误差，建议人工核对。")
            for run in p.runs:
                _set_run_font(run)
        if record.notes.strip():
            _add_label_value(doc, "核对说明：", record.notes.strip())
        if vr.issues:
            _add_label_value(doc, "自动校验提示：", "；".join(vr.issues))

        if idx != len(accepted):
            doc.add_paragraph("—" * 20)

    if review:
        doc.add_page_break()
        h = doc.add_paragraph(style="Heading 1")
        rr = h.add_run("需要人工核对的公式")
        _set_run_font(rr, size=16, bold=True)
        accepted_index = {c.candidate_id: i for i, (c, _, _) in enumerate(accepted, start=1)}
        for candidate, record, vr in review:
            p = doc.add_paragraph(style=None)
            p.style = doc.styles["Normal"]
            rr = p.add_run(f"公式 {accepted_index[candidate.candidate_id]} —— 第 {candidate.page_number} 页")
            _set_run_font(rr)
            if vr.issues:
                rr = p.add_run("：" + "；".join(vr.issues))
                _set_run_font(rr)

    if not accepted:
        p = doc.add_paragraph("未确认到可输出的数学公式。")
        for run in p.runs:
            _set_run_font(run)

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_docx)
    return output_docx
